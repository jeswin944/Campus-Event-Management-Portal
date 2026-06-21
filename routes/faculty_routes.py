from flask import Blueprint, render_template, request, redirect, url_for, session, flash, abort, make_response
from models.db import get_db_connection
from utils.helpers import login_required, role_required, add_notification, notify_admins
import openpyxl
import io

faculty_bp = Blueprint('faculty', __name__)

@faculty_bp.route('/faculty/dashboard')
@login_required
@role_required('faculty')
def faculty_dashboard():
    db = get_db_connection()
    cursor = db.cursor(dictionary=True)
    
    cursor.execute("""
        SELECT e.*, 
        COUNT(r.registration_id) as total_reg_count,
        SUM(CASE WHEN r.attendance = 'Present' THEN 1 ELSE 0 END) as attended_count
        FROM events e
        LEFT JOIN registrations r ON e.event_id = r.event_id
        WHERE e.coordinator_id = %s
        GROUP BY e.event_id
        ORDER BY e.event_date
    """, (session['user_id'],))
    events = cursor.fetchall()
    
    for e in events:
        total = e['total_reg_count']
        attended = e['attended_count'] or 0
        e['attendance_percentage'] = round((attended / total) * 100) if total > 0 else 0
    
    total_events = len(events)
    total_registrations = sum(e['total_reg_count'] for e in events)
    total_attendance = sum(e['attended_count'] for e in events if e['attended_count'])
    
    attendance_rate = 0
    if total_registrations > 0:
        attendance_rate = round((total_attendance / total_registrations) * 100, 1)

    analytics_data = []
    for e in events:
        analytics_data.append({
            'event_name': e['event_name'],
            'total_reg': e['total_reg_count'],
            'attended': int(e['attended_count']) if e['attended_count'] else 0
        })
    
    db.close()
    return render_template('faculty_dashboard.html', 
                           events=events,
                           total_events=total_events,
                           total_registrations=total_registrations,
                           attendance_rate=attendance_rate,
                           analytics_data=analytics_data)

@faculty_bp.route('/export-attendance/<int:event_id>')
@login_required
@role_required('faculty')
def export_attendance(event_id):
    db = get_db_connection()
    cursor = db.cursor(dictionary=True)
    
    cursor.execute("SELECT event_name FROM events WHERE event_id=%s AND coordinator_id=%s", (event_id, session['user_id']))
    event = cursor.fetchone()
    if not event:
        db.close()
        abort(403)
        
    cursor.execute("""
        SELECT s.name, s.register_number, s.department, s.email, s.semester, r.attendance, r.certificate_status
        FROM registrations r
        JOIN student s ON r.student_id = s.student_id
        WHERE r.event_id = %s
        ORDER BY s.name
    """, (event_id,))
    attendees = cursor.fetchall()
    db.close()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Sheet"
    
    headers = ["Name", "Register Number", "Department", "Semester", "Email", "Attendance Status", "Certificate Status"]
    ws.append(headers)
    
    for person in attendees:
        attendance = person['attendance'] if person['attendance'] else 'Absent'
        cert_status = person['certificate_status'] if person['certificate_status'] else 'Not Issued'
        ws.append([
            person['name'],
            person['register_number'],
            person['department'],
            person['semester'],
            person['email'],
            attendance,
            cert_status
        ])
        
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return make_response(output.getvalue(), 200, {
        "Content-Disposition": f"attachment; filename=Attendance_{event['event_name']}.xlsx",
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    })

@faculty_bp.route('/scan-attendance', methods=['GET', 'POST'])
@faculty_bp.route('/scan-attendance/<int:event_id>', methods=['GET', 'POST'])
@login_required
@role_required('faculty')
def scan_attendance(event_id=None):
    if request.method == 'POST':
        qr_token = request.form.get('qr_token')
        target_event_id = event_id or request.form.get('event_id')
        
        if not qr_token:
            return "No token provided", 400

        db = get_db_connection()
        cursor = db.cursor(dictionary=True)
        
        query = """
            SELECT r.registration_id, r.event_id, s.student_id, s.name, e.event_name 
            FROM registrations r 
            JOIN student s ON r.student_id = s.student_id 
            JOIN events e ON r.event_id = e.event_id 
            WHERE r.qr_token = %s
        """
        params = [qr_token]
        
        if target_event_id:
            query += " AND r.event_id = %s"
            params.append(target_event_id)
            
        cursor.execute(query, params)
        registration = cursor.fetchone()

        if registration:
            cursor.execute("""
                UPDATE registrations 
                SET attendance = 'Present', certificate_status = 'Pending' 
                WHERE qr_token = %s
            """, (qr_token,))
            db.commit()

            add_notification(registration['student_id'], 'student', f"Attendance marked: {registration['event_name']}.")
            notify_admins(f"Certificate Pending Approval: {registration['name']} - {registration['event_name']}.")
            
            db.close()
            return f"Success: Attendance marked for {registration['name']} (Event: {registration['event_name']})"
        else:
            db.close()
            if target_event_id:
                return "Error: Student not registered for THIS event, or Invalid QR Code", 404
            return "Error: Invalid QR Code", 404

    return render_template('scan_attendance.html', target_event_id=event_id)

@faculty_bp.route('/faculty/timetable')
@login_required
@role_required('faculty')
def faculty_timetable():
    db = get_db_connection()
    cursor = db.cursor(dictionary=True)
    
    cursor.execute("SELECT faculty_id, name, department FROM faculty WHERE faculty_id = %s", (session['user_id'],))
    faculty_info = cursor.fetchone()

    cursor.execute("""
        SELECT t.*, c.course_name 
        FROM timetable t 
        LEFT JOIN courses c ON t.course_id = c.course_id 
        WHERE t.faculty_id = %s 
        ORDER BY FIELD(day, 'Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday'), CAST(start_time AS TIME)
    """, (session['user_id'],))
    timetable = cursor.fetchall()
    
    db.close()
    return render_template('faculty_timetable.html', timetable=timetable, faculty=faculty_info)

@faculty_bp.route('/faculty/attendance')
@login_required
@role_required('faculty')
def daily_attendance_overview():
    db = get_db_connection()
    cursor = db.cursor(dictionary=True)

    from datetime import datetime
    now = datetime.now()
    today_name = now.strftime('%A')        # e.g. 'Thursday'
    today_date = now.strftime('%Y-%m-%d')  # e.g. '2026-02-27'

    # Optional date filter for history tab
    filter_date = request.args.get('filter_date', '').strip() or None

    # All slots for this faculty (full week)
    cursor.execute("""
        SELECT t.*, c.course_name
        FROM timetable t
        LEFT JOIN courses c ON t.course_id = c.course_id
        WHERE t.faculty_id = %s
        ORDER BY FIELD(day, 'Monday','Tuesday','Wednesday','Thursday','Friday','Saturday'), start_time
    """, (session['user_id'],))
    all_slots = cursor.fetchall()

    # Only today's slots (for the Today tab)
    today_slots = [s for s in all_slots if s['day'] == today_name]

    # ── Attendance History records (admin-style list) ──────────────────────────
    history_query = """
        SELECT
            da.attendance_date,
            da.timetable_id,
            c.department,
            c.semester,
            c.course_name,
            SUM(CASE WHEN da.status = 'Present' THEN 1 ELSE 0 END) AS present_count,
            COUNT(*) AS total_students
        FROM daily_attendance da
        JOIN timetable t ON da.timetable_id = t.timetable_id
        LEFT JOIN courses c ON t.course_id = c.course_id
        WHERE t.faculty_id = %s
    """
    params = [session['user_id']]
    if filter_date:
        history_query += " AND da.attendance_date = %s"
        params.append(filter_date)
    history_query += " GROUP BY da.attendance_date, da.timetable_id ORDER BY da.attendance_date DESC, da.timetable_id"
    cursor.execute(history_query, params)
    history_records = cursor.fetchall()

    db.close()
    return render_template('faculty_attendance.html',
                           all_slots=all_slots,
                           today_slots=today_slots,
                           today_name=today_name,
                           today_date=today_date,
                           history_records=history_records,
                           filter_date=filter_date)

@faculty_bp.route('/faculty/attendance-calendar')
@login_required
@role_required('faculty')
def attendance_calendar_api():
    """Return JSON attendance summary per date for FullCalendar."""
    from datetime import datetime
    from flask import jsonify
    db = get_db_connection()
    cursor = db.cursor(dictionary=True)

    # Get all timetable slots for this faculty
    cursor.execute("""
        SELECT t.timetable_id, c.course_name, t.department, t.semester, t.start_time, t.end_time
        FROM timetable t
        LEFT JOIN courses c ON t.course_id = c.course_id
        WHERE t.faculty_id = %s
    """, (session['user_id'],))
    slots = {s['timetable_id']: s for s in cursor.fetchall()}

    if not slots:
        db.close()
        return jsonify([])

    slot_ids = list(slots.keys())
    placeholders = ','.join(['%s'] * len(slot_ids))

    # Aggregate attendance per date per timetable slot
    cursor.execute(f"""
        SELECT da.attendance_date,
               da.timetable_id,
               SUM(CASE WHEN da.status='Present' THEN 1 ELSE 0 END) AS present_count,
               COUNT(*) AS total_count
        FROM daily_attendance da
        WHERE da.timetable_id IN ({placeholders})
        GROUP BY da.attendance_date, da.timetable_id
        ORDER BY da.attendance_date DESC
    """, slot_ids)
    rows = cursor.fetchall()
    db.close()

    # Group by date
    from collections import defaultdict
    date_map = defaultdict(lambda: {'present': 0, 'total': 0, 'classes': []})
    for row in rows:
        d = str(row['attendance_date'])
        date_map[d]['present'] += int(row['present_count'])
        date_map[d]['total']   += int(row['total_count'])
        slot = slots.get(row['timetable_id'], {})
        # Format time nicely
        def _fmt(t):
            if t is None: return ''
            if hasattr(t, 'seconds'):
                h, rem = divmod(int(t.total_seconds()), 3600)
                m = rem // 60
            else:
                parts = str(t).split(':')
                h, m = int(parts[0]), int(parts[1])
            h12 = h % 12 or 12
            return f"{h12}:{m:02d}"
        date_map[d]['classes'].append({
            'course': slot.get('course_name', 'Class'),
            'time': f"{_fmt(slot.get('start_time'))}–{_fmt(slot.get('end_time'))}",
            'present': int(row['present_count']),
            'total': int(row['total_count']),
        })

    # Build FullCalendar events
    events = []
    for date, info in date_map.items():
        p, t = info['present'], info['total']
        rate = round((p / t) * 100) if t else 0
        if rate >= 80:
            color = '#198754'   # green
        elif rate >= 50:
            color = '#fd7e14'   # orange
        else:
            color = '#dc3545'   # red
        events.append({
            'id': date,
            'title': f"{p}/{t} Present",
            'start': date,
            'backgroundColor': color,
            'borderColor': color,
            'textColor': '#fff',
            'extendedProps': {
                'present': p,
                'total': t,
                'rate': rate,
                'classes': info['classes'],
            }
        })
    return jsonify(events)

@faculty_bp.route('/faculty/mark-attendance/<int:timetable_id>', methods=['GET', 'POST'])
@login_required
@role_required('faculty')
def mark_students_attendance(timetable_id):
    db = get_db_connection()
    cursor = db.cursor(dictionary=True)
    
    from datetime import datetime
    selected_date_str = request.args.get('date') or request.form.get('attendance_date')
    if selected_date_str:
        try:
            selected_date = datetime.strptime(selected_date_str, '%Y-%m-%d').date()
        except:
            selected_date = datetime.now().date()
    else:
        selected_date = datetime.now().date()

    # Verify slot belongs to faculty
    cursor.execute("""
        SELECT t.*, c.course_name 
        FROM timetable t 
        LEFT JOIN courses c ON t.course_id = c.course_id 
        WHERE t.timetable_id = %s AND t.faculty_id = %s
    """, (timetable_id, session['user_id']))
    slot = cursor.fetchone()
    
    if not slot:
        db.close()
        abort(403)
        
    if request.method == 'POST':
        present_students = request.form.getlist('students') 
        
        cursor.execute("""
            SELECT student_id FROM student
            WHERE department = %s AND semester = %s
            AND register_number LIKE 'NIE23CS%'
        """, (slot['department'], slot['semester']))
        all_students = [s['student_id'] for s in cursor.fetchall()]

        # OD-approved students for this slot+date — faculty CANNOT override these
        cursor.execute("""
            SELECT student_id FROM onduty_requests
            WHERE timetable_id = %s AND od_date = %s AND status = 'Approved'
        """, (timetable_id, selected_date))
        od_protected = {row['student_id'] for row in cursor.fetchall()}
        
        try:
            for student_id in all_students:
                if student_id in od_protected:
                    continue   # Never override OD-approved Present
                status = 'Present' if str(student_id) in present_students else 'Absent'
                cursor.execute("""
                    INSERT INTO daily_attendance (timetable_id, student_id, attendance_date, status)
                    VALUES (%s, %s, %s, %s)
                    ON DUPLICATE KEY UPDATE status = VALUES(status)
                """, (timetable_id, student_id, selected_date, status))
            
            db.commit()
            flash(f"Attendance for {selected_date} marked successfully!", "success")
        except Exception as e:
            db.rollback()
            flash(f"Error marking attendance: {e}", "error")
            
        db.close()
        return redirect(url_for('faculty.daily_attendance_overview'))

    # GET: Fetch students for this slot
    cursor.execute("""
        SELECT * FROM student
        WHERE department = %s AND semester = %s
        AND register_number LIKE 'NIE23CS%'
        ORDER BY name
    """, (slot['department'], slot['semester']))
    students = cursor.fetchall()
    
    # Check if attendance already marked for selected date
    cursor.execute("SELECT student_id, status FROM daily_attendance WHERE timetable_id = %s AND attendance_date = %s", (timetable_id, selected_date))
    existing_attendance = {row['student_id']: row['status'] for row in cursor.fetchall()}

    # Find students with approved Duty Leave for this slot + date (their attendance is locked)
    cursor.execute("""
        SELECT student_id FROM onduty_requests
        WHERE timetable_id = %s AND od_date = %s AND status = 'Approved'
    """, (timetable_id, selected_date))
    od_student_ids = {row['student_id'] for row in cursor.fetchall()}
    
    db.close()
    return render_template('mark_attendance.html', slot=slot, students=students,
                           existing_attendance=existing_attendance,
                           today_date=selected_date,
                           od_student_ids=od_student_ids)

@faculty_bp.route('/faculty/manage-attendance/<int:event_id>', methods=['GET', 'POST'])
@login_required
@role_required('faculty')
def manage_attendance(event_id):
    db = get_db_connection()
    cursor = db.cursor(dictionary=True)
    
    # Verify event and faculty
    if session.get('is_admin'):
        cursor.execute("SELECT * FROM events WHERE event_id = %s", (event_id,))
    else:
        cursor.execute("SELECT * FROM events WHERE event_id = %s AND coordinator_id = %s", (event_id, session['user_id']))
    
    event = cursor.fetchone()
    if not event:
        db.close()
        abort(403)
        
    if request.method == 'POST':
        present_student_ids = request.form.getlist('attendance')
        # Reset all to Absent then mark present
        cursor.execute("UPDATE registrations SET attendance = 'Absent' WHERE event_id = %s", (event_id,))
        for sid in present_student_ids:
            cursor.execute("UPDATE registrations SET attendance = 'Present' WHERE event_id = %s AND student_id = %s", (event_id, sid))
        db.commit()
        flash(f"Attendance for {event['event_name']} updated.", "success")
        db.close()
        return redirect(url_for('faculty.manage_attendance', event_id=event_id))

    cursor.execute("""
        SELECT r.*, s.name, s.register_number, s.department 
        FROM registrations r
        JOIN student s ON r.student_id = s.student_id
        WHERE r.event_id = %s
        ORDER BY s.name
    """, (event_id,))
    registrations = cursor.fetchall()
    db.close()
    return render_template('manage_event_attendance.html', event=event, registrations=registrations)
